#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
================================================================================
CONECTA SUS - PROCESSADOR DE RELATÓRIOS LABORATORIAIS (CLI / DRAG & DROP)
================================================================================
Script de automação para extrair, filtrar e unificar dados de pacientes a partir
de relatórios laboratoriais (PDF) do Conecta SUS / Laboratório Municipal de Saúde.
================================================================================
"""

import sys
import os
import re
import unicodedata
from datetime import datetime
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import pdfplumber


def remover_acentos(texto: str) -> str:
    """
    Remove acentuação, caracteres especiais e converte para MAIÚSCULAS.
    Garante comparação resiliente contra variações de digitação.
    """
    if not texto:
        return ""
    nfkd = unicodedata.normalize('NFKD', str(texto))
    ascii_texto = nfkd.encode('ASCII', 'ignore').decode('utf-8')
    # Substitui múltiplos espaços por um só e limpa pontuações
    limpo = re.sub(r'[^a-zA-Z0-9\s]', ' ', ascii_texto.upper())
    return " ".join(limpo.split())


def limpar_celula(valor) -> str:
    """
    Remove quebras de linha (\n, \r) e espaços extras de células da tabela.
    """
    if valor is None:
        return ""
    texto = str(valor)
    texto = texto.replace('\r\n', ' ').replace('\n', ' ').replace('\r', ' ')
    return " ".join(texto.split()).strip()


def identificar_medico(medico_raw: str) -> str | None:
    """
    Identifica se o médico corresponde a:
      - Dr. Douglas Otavio J. Junior (palavras-chave: DOUGLAS, DOUGLAS J. JR, DOUGLAS OTAVIO)
      - Dra. Livia Zavao Zanetti (variações: LIVIA, ZAVAO, ZAVON, ZANETTI)
    
    Retorna o nome canônico ou None se for outro médico (a ser ignorado).
    """
    if not medico_raw:
        return None
    
    norm = remover_acentos(medico_raw)
    
    # 1. Correspondência para Dr. Douglas
    if "DOUGLAS" in norm and ("OTAVIO" in norm or "JR" in norm or "JUNIOR" in norm or len(norm.split()) <= 2):
        return "Dr. Douglas Otavio J. Junior"
        
    # 2. Correspondência para Dra. Lívia
    if "LIVIA" in norm or "ZAVAO" in norm or "ZAVON" in norm or "ZANETTI" in norm:
        return "Dra. Livia Zavao Zanetti"
        
    return None


def mapear_colunas(cabecalho: list[str]) -> dict[str, int]:
    """
    Mapeia dinamicamente os índices das colunas a partir dos termos do cabeçalho.
    """
    cabecalho_norm = [remover_acentos(c) for c in cabecalho]
    
    sinonimos = {
        'data': ['DATA COLETA', 'DT COLETA', 'DATA', 'DT AGENDAMENTO', 'AGENDAMENTO', 'DATA ATENDIMENTO', 'COLETA'],
        'prontuario': ['PRONTUARIO', 'PRONT', 'N PRONTUARIO', 'N PRONT', 'PRONT NO', 'CNS', 'CARTAO SUS', 'CODIGO'],
        'paciente': ['NOME DO PACIENTE', 'PACIENTE', 'NOME', 'USUARIO', 'NOME PACIENTE'],
        'idade': ['IDADE', 'ID', 'ANOS', 'FAIXA ETARIA', 'DT NASC'],
        'medico': ['MEDICO SOLICITANTE', 'MEDICO', 'SOLICITANTE', 'PROFISSIONAL', 'NOME DO MEDICO', 'PRESCRITOR']
    }
    
    mapeamento = {}
    for chave, lista_sinonimos in sinonimos.items():
        idx_encontrado = -1
        # Correspondência exata
        for sin in lista_sinonimos:
            for i, col in enumerate(cabecalho_norm):
                if sin == col:
                    idx_encontrado = i
                    break
            if idx_encontrado != -1:
                break
        
        # Substring
        if idx_encontrado == -1:
            for sin in lista_sinonimos:
                for i, col in enumerate(cabecalho_norm):
                    if sin in col:
                        idx_encontrado = i
                        break
                if idx_encontrado != -1:
                    break
                    
        mapeamento[chave] = idx_encontrado
        
    return mapeamento


def extrair_dados_pdf(caminho_pdf: str) -> list[dict]:
    """
    Lê todas as páginas do relatório em PDF usando pdfplumber e extrai as linhas das tabelas.
    """
    registros = []
    
    with pdfplumber.open(caminho_pdf) as pdf:
        total_paginas = len(pdf.pages)
        print(f"[*] Analisando {total_paginas} pagina(s) do documento...")
        
        mapeamento_ativo = None
        
        for num_pag, pagina in enumerate(pdf.pages, 1):
            tabelas = pagina.extract_tables()
            
            # Se extract_tables não retornar lista, tenta extract_table isolado
            if not tabelas:
                tabela_unica = pagina.extract_table()
                if tabela_unica:
                    tabelas = [tabela_unica]
            
            if not tabelas:
                continue
                
            for tabela in tabelas:
                if not tabela or len(tabela) < 1:
                    continue
                    
                inicio_dados = 0
                candidato_cabecalho = [limpar_celula(cel) for cel in tabela[0]]
                mapa_teste = mapear_colunas(candidato_cabecalho)
                
                # Se identificou paciente ou médico, a linha 0 é cabeçalho
                if mapa_teste.get('paciente', -1) != -1 or mapa_teste.get('medico', -1) != -1:
                    mapeamento_ativo = mapa_teste
                    inicio_dados = 1
                elif mapeamento_ativo is None:
                    # Layout padrão SUS de 5/6 colunas caso não haja texto claro de cabeçalho
                    mapeamento_ativo = {
                        'data': 0 if len(tabela[0]) > 0 else -1,
                        'prontuario': 1 if len(tabela[0]) > 1 else -1,
                        'paciente': 2 if len(tabela[0]) > 2 else -1,
                        'idade': 3 if len(tabela[0]) > 3 else -1,
                        'medico': 4 if len(tabela[0]) > 4 else -1,
                    }
                
                for linha in tabela[inicio_dados:]:
                    if not linha:
                        continue
                        
                    def obter_campo(chave):
                        idx = mapeamento_ativo.get(chave, -1)
                        if idx != -1 and idx < len(linha):
                            return limpar_celula(linha[idx])
                        return ""
                    
                    paciente = obter_campo('paciente')
                    medico = obter_campo('medico')
                    
                    # Ignora repetições acidentais de cabeçalho ou linhas em branco
                    if not paciente or remover_acentos(paciente) in ['NOME', 'PACIENTE', 'NOME DO PACIENTE']:
                        continue
                        
                    registros.append({
                        'data': obter_campo('data'),
                        'prontuario': obter_campo('prontuario'),
                        'paciente': paciente,
                        'idade': obter_campo('idade'),
                        'medico': medico
                    })
                    
    return registros


def processar_e_deduplicar(registros_brutos: list[dict]) -> tuple[list[dict], int, int]:
    """
    1. Filtra apenas os médicos solicitantes alvo (Dr. Douglas e Dra. Lívia).
    2. Deduplica pacientes por ['Prontuário', 'Paciente'] (cada pessoa aparece uma vez).
    3. Unifica solicitações conjuntas (se houver Dr. Douglas e Dra. Lívia, exibe ambos).
    4. Ordena por Médico Solicitante e Nome do Paciente.
    """
    total_lidos = len(registros_brutos)
    filtrados_medico = []
    
    # 1. Filtro de Médicos
    for reg in registros_brutos:
        medico_canonico = identificar_medico(reg.get('medico', ''))
        if medico_canonico:
            item = dict(reg)
            item['medico_canonico'] = medico_canonico
            filtrados_medico.append(item)
            
    total_filtrados = len(filtrados_medico)
    
    # 2. Agrupamento e Deduplicação por ['Prontuário', 'Paciente']
    agrupados = {}
    
    for reg in filtrados_medico:
        prontuario = reg.get('prontuario', '').strip()
        paciente = reg.get('paciente', '').strip()
        paciente_norm = remover_acentos(paciente)
        
        # Chave de unificação por prontuário e nome
        chave = f"{prontuario}|{paciente_norm}" if prontuario else f"SEM_PRONT|{paciente_norm}"
        
        if chave not in agrupados:
            agrupados[chave] = {
                'data': reg.get('data', '').strip(),
                'prontuario': prontuario,
                'paciente': paciente,
                'idade': reg.get('idade', '').strip(),
                'medicos': {reg['medico_canonico']}
            }
        else:
            # Adiciona médico caso seja diferente
            agrupados[chave]['medicos'].add(reg['medico_canonico'])
            # Preenche idade ou data se faltavam anteriormente
            if not agrupados[chave]['idade'] and reg.get('idade'):
                agrupados[chave]['idade'] = reg.get('idade').strip()
            if not agrupados[chave]['data'] and reg.get('data'):
                agrupados[chave]['data'] = reg.get('data').strip()
                
    # 3. Consolidação dos dados finais
    consolidados = []
    for item in agrupados.values():
        medicos_set = item['medicos']
        # Se possuir pedidos de ambos os médicos:
        if len(medicos_set) > 1:
            medico_final = "Dr. Douglas Otavio J. Junior / Dra. Livia Zavao Zanetti"
        else:
            medico_final = next(iter(medicos_set))
            
        consolidados.append({
            'Data da Coleta': item['data'],
            'Prontuário': item['prontuario'],
            'Nome do Paciente': item['paciente'],
            'Idade': item['idade'],
            'Médico Solicitante': medico_final
        })
        
    # 4. Ordenação por Médico Solicitante e Nome do Paciente (alfabética)
    consolidados.sort(key=lambda x: (x['Médico Solicitante'], remover_acentos(x['Nome do Paciente'])))
    
    return consolidados, total_lidos, total_filtrados


def exportar_para_excel(dados: list[dict], caminho_saida: str):
    """
    Gera um arquivo Excel (.xlsx) com auto-ajuste de colunas e formatação visual limpa.
    """
    df = pd.DataFrame(dados)
    
    with pd.ExcelWriter(caminho_saida, engine='openpyxl') as writer:
        df.to_excel(writer, sheet_name='Pacientes Filtrados', index=False, startrow=2)
        workbook = writer.book
        worksheet = writer.sheets['Pacientes Filtrados']
        
        # Cores e Estilos
        azul_cabecalho = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
        fonte_cabecalho = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        
        zebrado_par = PatternFill(start_color="F2F5F8", end_color="F2F5F8", fill_type="solid")
        zebrado_impar = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
        
        borda_fina = Border(
            left=Side(style='thin', color='D9D9D9'),
            right=Side(style='thin', color='D9D9D9'),
            top=Side(style='thin', color='D9D9D9'),
            bottom=Side(style='thin', color='D9D9D9')
        )
        
        # Título superior
        worksheet.merge_cells("A1:E1")
        cel_titulo = worksheet["A1"]
        cel_titulo.value = "RELATÓRIO CONSOLIDADO DE COLETA LABORATORIAL - CONECTA SUS"
        cel_titulo.font = Font(name="Calibri", size=13, bold=True, color="1F4E79")
        cel_titulo.alignment = Alignment(horizontal="left", vertical="center")
        worksheet.row_dimensions[1].height = 24
        worksheet.row_dimensions[3].height = 24
        
        # Cabeçalho da tabela (linha 3)
        for col_idx in range(1, len(df.columns) + 1):
            celula = worksheet.cell(row=3, column=col_idx)
            celula.fill = azul_cabecalho
            celula.font = fonte_cabecalho
            celula.alignment = Alignment(horizontal="center", vertical="center")
            celula.border = borda_fina
            
        # Linhas de dados (linha 4 em diante)
        for row_idx in range(4, len(df) + 4):
            worksheet.row_dimensions[row_idx].height = 20
            preenchimento = zebrado_par if (row_idx % 2 == 0) else zebrado_impar
            
            for col_idx in range(1, len(df.columns) + 1):
                celula = worksheet.cell(row=row_idx, column=col_idx)
                celula.fill = preenchimento
                celula.border = borda_fina
                celula.font = Font(name="Calibri", size=10)
                
                nome_coluna = df.columns[col_idx - 1]
                if nome_coluna in ['Data da Coleta', 'Prontuário', 'Idade']:
                    celula.alignment = Alignment(horizontal="center", vertical="center")
                    if nome_coluna == 'Prontuário':
                        celula.number_format = '@'  # Preserva zeros à esquerda
                else:
                    celula.alignment = Alignment(horizontal="left", vertical="center")
                    
        # Auto-ajuste da largura das colunas
        for col in worksheet.columns:
            maior_tamanho = 0
            letra_col = get_column_letter(col[0].column)
            for celula in col:
                if celula.row < 3:
                    continue
                conteudo = str(celula.value or '')
                maior_tamanho = max(maior_tamanho, len(conteudo))
            worksheet.column_dimensions[letra_col].width = max(maior_tamanho + 5, 14)


def main():
    print("=" * 70)
    print("   CONECTA SUS - EXTRATOR E FILTRO DE LAUDOS LABORATORIAIS")
    print("=" * 70)
    
    # 1. Validação do argumento de entrada
    if len(sys.argv) < 2:
        print("\n[ERRO] Nenhum arquivo foi informado!")
        print("Uso: python processador.py \"caminho\\para\\relatorio.pdf\"")
        print("Ou simplesmente arraste o PDF para cima do arquivo FILTRAR_PDF.bat.\n")
        sys.exit(1)
        
    caminho_pdf = sys.argv[1].strip()
    # Remove aspas extras se presentes
    if caminho_pdf.startswith('"') and caminho_pdf.endswith('"'):
        caminho_pdf = caminho_pdf[1:-1]
        
    if not os.path.exists(caminho_pdf):
        print(f"\n[ERRO] O arquivo especificado nao foi encontrado:\n{caminho_pdf}\n")
        sys.exit(1)
        
    if not caminho_pdf.lower().endswith('.pdf'):
        print("\n[ERRO] O arquivo informado nao e um documento PDF (.pdf).\n")
        sys.exit(1)
        
    nome_arquivo = os.path.basename(caminho_pdf)
    pasta_destino = os.path.dirname(os.path.abspath(caminho_pdf))
    nome_base = os.path.splitext(nome_arquivo)[0]
    caminho_excel = os.path.join(pasta_destino, f"{nome_base}_FILTRADO.xlsx")
    
    print(f"\n[1/3] Arquivo recebido: {nome_arquivo}")
    print(f"      Pasta: {pasta_destino}")
    
    # 2. Leitura do PDF
    try:
        registros_brutos = extrair_dados_pdf(caminho_pdf)
    except Exception as e:
        print(f"\n[FALHA] Nao foi possivel ler o arquivo PDF: {e}")
        sys.exit(1)
        
    if not registros_brutos:
        print("\n[AVISO] Nenhuma tabela ou linha de agendamento foi identificada no PDF.")
        sys.exit(1)
        
    print(f"[2/3] Registros extraidos do PDF: {len(registros_brutos)}")
    
    # 3. Filtragem e Deduplicação
    consolidados, total_lidos, total_filtrados = processar_e_deduplicar(registros_brutos)
    
    if not consolidados:
        print("\n[AVISO] Nenhum paciente encontrado para os medicos alvo:")
        print("        - Dr. Douglas Otavio J. Junior")
        print("        - Dra. Livia Zavao Zanetti")
        print(f"        (Foram lidos {total_lidos} registros de outros medicos).\n")
        sys.exit(0)
        
    print(f"      -> Registros correspondentes aos medicos alvo: {total_filtrados}")
    print(f"      -> Total de pacientes unicos apos deduplicacao: {len(consolidados)}")
    
    # 4. Exportação para Excel
    print("\n[3/3] Gerando planilha Excel formatada...")
    try:
        exportar_para_excel(consolidados, caminho_excel)
    except Exception as e:
        print(f"\n[FALHA] Erro ao salvar a planilha Excel: {e}")
        sys.exit(1)
        
    print("=" * 70)
    print("   PROCESSAMENTO CONCLUIDO COM SUCESSO!")
    print("=" * 70)
    print(f"Arquivo gerado: {os.path.basename(caminho_excel)}")
    print(f"Localizacao:    {caminho_excel}\n")
    

if __name__ == "__main__":
    main()
